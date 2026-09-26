"""Unit tests for document ingestion service (PDF, DOCX, TXT)."""
import io

from docx import Document
from pypdf import PdfWriter

from app.services.ingest import extract_text_from_bytes, extract_text_from_raw_string


def test_extract_from_plain_text_bytes():
    raw_bytes = b"Section 1: Payment\n\nSection 2: Term"
    pages = extract_text_from_bytes("test.txt", raw_bytes)
    assert len(pages) == 1
    assert "Section 1" in pages[0][1]


def test_extract_from_docx_bytes():
    doc = Document()
    doc.add_paragraph("Paragraph 1 in DOCX document.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Cell A"
    table.rows[0].cells[1].text = "Cell B"

    bio = io.BytesIO()
    doc.save(bio)
    docx_bytes = bio.getvalue()

    pages = extract_text_from_bytes("contract.docx", docx_bytes)
    assert len(pages) == 1
    assert "Paragraph 1 in DOCX" in pages[0][1]
    assert "Cell A | Cell B" in pages[0][1]


def test_extract_from_pdf_bytes():
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    bio = io.BytesIO()
    writer.write(bio)
    pdf_bytes = bio.getvalue()

    pages = extract_text_from_bytes("lease.pdf", pdf_bytes)
    # Blank page returns empty text list
    assert isinstance(pages, list)


def test_extract_from_raw_string():
    raw = "Page 1 content\fPage 2 content"
    pages = extract_text_from_raw_string(raw)
    assert len(pages) == 2
    assert pages[0][0] == 1
    assert pages[1][0] == 2
